"""Synthetic upload tests; does not read user files or print extracted content."""
import argparse, base64, concurrent.futures, io, json, pathlib, re, time
import urllib.error, urllib.parse, urllib.request
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

ENDPOINT = 'https://marindsain8n.ru/webhook/ekt-attachment-extract'
TEXT = 'SKU 200300285_ | Legrand DPX3 3P 160A | Quantity 2'


def fixture_font():
    for name in ('C:/Windows/Fonts/arial.ttf',
                 '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
                 '/usr/share/fonts/dejavu/DejaVuSans.ttf', 'DejaVuSans.ttf'):
        try:
            return ImageFont.truetype(name, 36)
        except OSError:
            pass
    try:
        return ImageFont.load_default(size=36)
    except TypeError:
        return ImageFont.load_default()


def fixtures():
    stream = io.BytesIO()
    pdf = canvas.Canvas(stream)
    pdf.drawString(60, 760, TEXT)
    pdf.save()
    yield 'synthetic.pdf', 'application/pdf', stream.getvalue()

    stream = io.BytesIO()
    doc = Document()
    doc.add_heading('Product specification', 0)
    doc.add_paragraph(TEXT)
    doc.save(stream)
    yield 'synthetic.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', stream.getvalue()

    stream = io.BytesIO()
    book = Workbook()
    book.active.append(['SKU', 'Product', 'Quantity'])
    book.active.append(['200300285_', 'Legrand DPX3 3P 160A', 2])
    book.save(stream)
    yield 'synthetic.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', stream.getvalue()

    stream = io.BytesIO()
    img = Image.new('RGB', (1600, 300), 'white')
    ImageDraw.Draw(img).text((40, 80), TEXT, fill='black', font=fixture_font())
    img.save(stream, format='JPEG')
    yield 'synthetic.jpg', 'image/jpeg', stream.getvalue()
    yield 'synthetic.csv', 'text/csv', b'sku,name,quantity\n200300285_,Legrand DPX3 3P 160A,2\n'


def post(fixture, endpoint=ENDPOINT):
    name, mime, data = fixture
    payload = {'fileName': name, 'mimeType': mime,
               'dataUrl': 'data:' + mime + ';base64,' + base64.b64encode(data).decode()}
    started = time.monotonic()
    req = urllib.request.Request(endpoint, json.dumps(payload).encode(), {
        'Content-Type': 'application/json',
        'Origin': 'https://ekt-assistant-hackalem-2026.koncevojdanila10.chatgpt.site',
    }, method='POST')
    status, cors, result, error = 0, None, {}, None
    try:
        try:
            response = urllib.request.urlopen(req, timeout=85)
        except urllib.error.HTTPError as exc:
            response = exc
        with response:
            status = response.status
            cors = response.headers.get('Access-Control-Allow-Origin')
            result = json.loads(response.read(65537))
        if not isinstance(result, dict):
            result, error = {}, 'invalid_response_shape'
    except (urllib.error.URLError, TimeoutError, OSError, ValueError) as exc:
        error = type(exc).__name__
    extracted = result.get('extracted_text', '')
    if not isinstance(extracted, str):
        extracted = ''
    sku_found = '200300285_' in extracted
    quantity_found = re.search(r'\b2\b', extracted) is not None
    passed = status == 200 and result.get('ok') is True and sku_found and quantity_found and len(extracted) <= 6000
    return {'file': name, 'status': status, 'pass': passed,
            'seconds': round(time.monotonic() - started, 2), 'cors': cors,
            'extracted_chars': len(extracted), 'sku_found': sku_found,
            'quantity_found': quantity_found, 'error_type': error}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--endpoint', default=ENDPOINT, help='HTTP(S) extraction endpoint; HTTP allowed for local Docker.')
    parser.add_argument('--write-fixtures', action='store_true', help='Save fixtures locally; no network requests.')
    args = parser.parse_args(argv)
    parsed = urllib.parse.urlparse(args.endpoint)
    if parsed.scheme not in ('http', 'https') or not parsed.netloc or parsed.username or parsed.password:
        parser.error('--endpoint must be HTTP(S) without embedded credentials')
    if args.write_fixtures:
        folder = pathlib.Path(__file__).parent / 'fixtures'
        folder.mkdir(exist_ok=True)
        for name, _, data in fixtures():
            (folder / name).write_bytes(data)
        print(str(folder.resolve()))
        return 0
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda fixture: post(fixture, args.endpoint), fixtures()))
    for result in results:
        print(json.dumps(result, ensure_ascii=False), flush=True)
    return 0 if all(result['pass'] for result in results) else 1


if __name__ == '__main__':
    raise SystemExit(main())
